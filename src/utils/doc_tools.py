from __future__ import annotations
from pptx import Presentation
from docx import Document
import io

def make_ppt(title="Client Update", bullets=None) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    if bullets:
        s2 = prs.slides.add_slide(prs.slide_layouts[1])
        s2.shapes.title.text = "Key Points"
        tf = s2.shapes.placeholders[1].text_frame
        tf.clear()
        for b in bullets:
            p = tf.add_paragraph()
            p.text = b
            p.level = 0
    bio = io.BytesIO()
    prs.save(bio)
    bio.seek(0)
    return bio.read()

def make_docx(title="Executive Summary", paragraphs=None) -> bytes:
    doc = Document()
    doc.add_heading(title, 0)
    if paragraphs:
        for p in paragraphs:
            doc.add_paragraph(p)
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.read()

def summarize_text(text: str, max_chars=1200) -> str:
    if not text:
        return ""
    text = text[:max_chars]
    return text
